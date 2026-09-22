"""Sacar tablas a Excel y a CSV.

Dos mitades. La de arriba es `api/tablas.py` a solas: qué celda se convierte en
número y cuál no, que es la decisión que puede estropear una hoja sin que se
note hasta que alguien ha trabajado sobre ella. La de abajo dibuja una factura
con rejilla y comprueba que lo que sale es lo que ponía, y que el recuadro con un
párrafo dentro —que también es una rejilla— no se cuela como tabla.
"""
import fitz
import pytest

from tests.conftest import SESION, subida


@pytest.fixture
def cliente(entorno):
    entorno()
    import app as modulo
    return modulo.create_app('pesados').test_client()


FACTURA = [
    ['Concepto', 'Uds', 'Precio', 'Importe'],
    ['Horas de desarrollo', '12', '45,00', '540,00'],
    ['Licencia anual', '1', '1.200,00', '1.200,00'],
    ['Descuento', '1', '(120,00)', '(120,00)'],
    ['Total', '', '', '1.620,00'],
]


def factura(con_recuadro=True, filas=FACTURA):
    """Un PDF con una tabla dibujada y, si se pide, un recuadro que no lo es."""
    archivo = fitz.open()
    pagina = archivo.new_page()
    anchos, alto, y0 = [150, 50, 90, 90], 24, 100
    for i, fila in enumerate(filas):
        x = 60
        for j, celda in enumerate(fila):
            rect = fitz.Rect(x, y0 + i * alto, x + anchos[j], y0 + (i + 1) * alto)
            pagina.draw_rect(rect, color=(0, 0, 0), width=0.7)
            pagina.insert_text((x + 4, y0 + i * alto + 16), celda, fontsize=9)
            x += anchos[j]
    if con_recuadro:
        pagina.draw_rect(fitz.Rect(60, 260, 400, 320), color=(0, 0, 0), width=0.7)
        pagina.insert_text((66, 280), 'Aviso: esto es un recuadro, no una tabla.', fontsize=9)
    datos = archivo.tobytes()
    archivo.close()
    return datos


def subir(datos, nombre='factura.pdf'):
    from storage import storage
    return storage.save_upload(SESION, subida(datos, nombre)).id


def extraer(cliente, datos, **opciones):
    return cliente.post('/api/tools/extraer-tablas', headers={'X-Session-Id': SESION},
                        json={'file_ids': [subir(datos)], **opciones})


def leer(respuesta, indice=0):
    from storage import storage
    return storage.path_of(SESION, respuesta.get_json()['files'][indice]['id'])


# --- qué es un número y qué es texto -------------------------------------

@pytest.mark.parametrize('texto, valor', [
    ('1.234,56', 1234.56),
    ('12,50', 12.5),
    ('-1.000', -1000.0),
    ('+3,5', 3.5),
    ('1.234,56 €', 1234.56),
    ('€ 12', 12.0),
    ('(120,00)', -120.0),       # el negativo de toda la vida en un balance
])
def test_un_importe_se_convierte_en_numero(texto, valor):
    from api.tablas import a_numero

    assert a_numero(texto) == pytest.approx(valor)


@pytest.mark.parametrize('texto', [
    '08001',        # un código postal convertido a 8001 ya no es un código postal
    '2026',
    '12345678',
    '45%',
    '1.234,56 %',   # convertirlo pierde el tanto por ciento; a 0,45, la cifra
    '(1.234,56',    # un paréntesis suelto es un error de lectura, no un negativo
    'Total',
    '',
    None,
])
def test_lo_que_puede_ser_otra_cosa_se_queda_como_texto(texto):
    from api.tablas import a_numero

    assert a_numero(texto) is None


def test_el_csv_lleva_bom_y_punto_y_coma():
    """Con comas y sin BOM, un Excel en español abre la hoja en una columna y
    con los acentos rotos, que es como se consigue que nadie vuelva a usarlo."""
    from api.tablas import a_csv

    salida = a_csv([['Año', 'Importe'], ['2026', '1.234,56']])

    assert salida.startswith(b'\xef\xbb\xbf')
    assert salida.decode('utf-8-sig').splitlines()[0] == 'Año;Importe'


def test_las_celdas_con_salto_de_linea_se_juntan():
    """El salto dentro de una celda es de la maquetación del PDF, no del dato."""
    from api.tablas import a_csv

    salida = a_csv([['Concepto\nmuy largo', '1']]).decode('utf-8-sig')

    assert salida.splitlines()[0] == 'Concepto muy largo;1'


def test_dos_hojas_no_pueden_llamarse_igual(tmp_path):
    from openpyxl import load_workbook

    from api.tablas import a_xlsx

    destino = str(tmp_path / 'libro.xlsx')
    a_xlsx([('Página 1', [['a']]), ('Página 1', [['b']]), ('Página 1', [['c']])], destino)

    assert load_workbook(destino).sheetnames == ['Página 1', 'Página 1 (2)', 'Página 1 (3)']


def test_un_nombre_de_hoja_no_puede_llevar_lo_que_excel_prohibe(tmp_path):
    from openpyxl import load_workbook

    from api.tablas import a_xlsx

    destino = str(tmp_path / 'libro.xlsx')
    a_xlsx([('Ventas/2026: [enero]', [['a']])], destino)

    assert load_workbook(destino).sheetnames == ['Ventas 2026   enero']


# --- la herramienta, leyendo lo que escribe ------------------------------

def test_la_factura_sale_con_sus_importes_como_numeros(cliente):
    from openpyxl import load_workbook

    respuesta = extraer(cliente, factura(), formato='xlsx')

    assert respuesta.status_code == 201
    libro = load_workbook(leer(respuesta))
    assert libro.sheetnames == ['Página 1']
    filas = list(libro['Página 1'].iter_rows(values_only=True))
    assert filas[0] == ('Concepto', 'Uds', 'Precio', 'Importe')
    assert filas[1][3] == pytest.approx(540.0)
    assert filas[3][3] == pytest.approx(-120.0)    # (120,00)
    assert filas[1][1] == '12'                     # unidades: un entero pelado, texto


def test_un_recuadro_con_un_parrafo_no_es_una_tabla(cliente):
    """`find_tables` marca cualquier rejilla, y en un PDF hay rejillas por todas
    partes. Entregar una tabla inventada es peor que no entregar ninguna."""
    respuesta = cliente.post('/api/tools/extraer-tablas/inspeccionar',
                             headers={'X-Session-Id': SESION},
                             json={'file_ids': [subir(factura())]})

    assert respuesta.get_json()['total'] == 1


def test_un_pdf_sin_tablas_lo_dice_en_vez_de_devolver_un_libro_vacio(cliente):
    archivo = fitz.open()
    archivo.new_page().insert_text((72, 100), 'Sólo texto corrido.', fontsize=12)
    datos = archivo.tobytes()
    archivo.close()

    respuesta = extraer(cliente, datos)

    assert respuesta.status_code == 422
    assert 'espacios' in respuesta.get_json()['error']


def test_en_csv_sale_un_archivo_por_tabla(cliente):
    respuesta = extraer(cliente, factura(), formato='csv')

    archivos = respuesta.get_json()['files']
    assert [a['name'] for a in archivos] == ['factura-pagina-1-tabla-1.csv']
    texto = open(leer(respuesta), 'rb').read().decode('utf-8-sig')
    assert texto.splitlines()[0] == 'Concepto;Uds;Precio;Importe'


def test_la_inspeccion_no_escribe_nada(cliente):
    respuesta = cliente.post('/api/tools/extraer-tablas/inspeccionar',
                             headers={'X-Session-Id': SESION},
                             json={'file_ids': [subir(factura())]})

    cuerpo = respuesta.get_json()
    assert cuerpo['tablas'] == [{'pagina': 1, 'filas': 5, 'columnas': 4}]


def test_lo_que_no_es_un_pdf_se_rechaza(cliente):
    from storage import storage

    ident = storage.save_upload(SESION, subida(b'a;b\n1;2\n', 'datos.csv')).id
    respuesta = cliente.post('/api/tools/extraer-tablas', headers={'X-Session-Id': SESION},
                             json={'file_ids': [ident]})

    assert respuesta.status_code == 400
    assert 'no es un PDF' in respuesta.get_json()['error']
